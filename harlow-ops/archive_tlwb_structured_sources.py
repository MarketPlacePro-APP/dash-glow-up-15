#!/usr/bin/env python3
"""Archive current TLWB structured exports into the canonical SQLite warehouse.

The dashboard continues to render static artifacts, but this loader makes the
same fetched XLSX/JSON source truth queryable for ad-hoc analysis. It is
append-only by content hash: unchanged source files create a new check record,
not duplicate raw rows.
"""
from __future__ import annotations

import hashlib
import json
import os
import sqlite3
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

ROOT = Path(os.environ.get("TLWB_WORKSPACE_ROOT", "/Users/seanwilliams/.openclaw/workspace-main"))
DB_PATH = Path(os.environ.get("TLWB_SOURCE_ARCHIVE_DB", "/Users/seanwilliams/.openclaw/workspace/outputs/tlwb_source_archive.db"))
TZ = ZoneInfo("America/Denver")


@dataclass(frozen=True)
class SourceSpec:
    source_id: str
    source_name: str
    source_type: str
    source_family: str
    business_area: str
    canonical_ref: str
    local_path: Path
    kind: str


SOURCES = (
    SourceSpec(
        "numbers_per_session",
        "Numbers Per Session Google Sheet",
        "google_sheet_export",
        "session_sheet",
        "TLWB",
        "https://docs.google.com/spreadsheets/d/1dfke_KCSGHNfnG_FUjAwo1TPAXFtgLEh9tE_XqQB0TU/edit?usp=sharing",
        ROOT / "data/google_exports/numbers_per_session_1dfke_latest.xlsx",
        "xlsx",
    ),
    SourceSpec(
        "upcoming_schedule",
        "Upcoming Schedule Google Sheet",
        "google_sheet_export",
        "schedule",
        "TLWB",
        "https://docs.google.com/spreadsheets/d/1F05mJPz4m8Kzxc8ROTQc4puRBky263ghSUMTg4kxKqY/edit?usp=sharing",
        ROOT / "data/google_exports/upcoming_schedule_1F05mJ_latest.xlsx",
        "xlsx",
    ),
    SourceSpec(
        "workshop_team_schedule",
        "Workshop Team Scheduling Google Sheet",
        "google_sheet_export",
        "schedule",
        "TLWB",
        "https://docs.google.com/spreadsheets/d/1psHz1be5AdbpjLu4vWEIvecf6CLeuRWodBoHu20Dotw/edit?usp=sharing",
        ROOT / "data/google_exports/workshop_schedule_sheet_1psHz1_latest.xlsx",
        "xlsx",
    ),
    SourceSpec(
        "lindsey_2026_ws_sales_tracker",
        "2026 WS Sales Tracker",
        "google_sheet_export",
        "lindsey",
        "TLWB",
        "https://docs.google.com/spreadsheets/d/1CmJYo4jIiweNArfZvvKdb0q_WlLtLsNH1UqHaad5gxQ/edit",
        ROOT / "data/google_exports/ws_sales_tracker_1CmJ_latest.xlsx",
        "xlsx",
    ),
    SourceSpec(
        "replit_inside_sales_dpl",
        "TLWB Inside Sales DPL API",
        "json_api_export",
        "inside_sales",
        "TLWB",
        "https://utltlwb-stats.replit.app/api/tlwb/inside-sales-dpl",
        ROOT / "data/inside_replit_latest.json",
        "inside_sales_json",
    ),
    SourceSpec(
        "replit_collections_performance",
        "TLWB Collections Performance API",
        "json_api_export",
        "collections",
        "TLWB",
        "https://utltlwb-stats.replit.app/api/tlwb/collections-performance",
        ROOT / "data/collections_replit_latest.json",
        "collections_json",
    ),
)


def now_iso() -> str:
    return datetime.now(TZ).isoformat()


def json_default(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=json_default)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def init_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        PRAGMA busy_timeout=30000;
        CREATE TABLE IF NOT EXISTS source_rows_raw (
          row_raw_id INTEGER PRIMARY KEY,
          snapshot_id TEXT NOT NULL,
          source_id TEXT NOT NULL,
          record_type TEXT NOT NULL,
          sheet_name TEXT,
          row_number INTEGER,
          record_key TEXT,
          row_json TEXT NOT NULL,
          row_sha256 TEXT NOT NULL,
          ingested_at TEXT NOT NULL,
          UNIQUE(snapshot_id, record_type, sheet_name, row_number, record_key)
        );
        CREATE INDEX IF NOT EXISTS idx_source_rows_raw_source_snapshot
          ON source_rows_raw(source_id, snapshot_id);
        CREATE INDEX IF NOT EXISTS idx_source_rows_raw_sheet
          ON source_rows_raw(source_id, sheet_name, row_number);

        CREATE TABLE IF NOT EXISTS inside_sales_dpl_fact (
          fact_id INTEGER PRIMARY KEY,
          snapshot_id TEXT NOT NULL,
          source_id TEXT NOT NULL,
          observed_at TEXT NOT NULL,
          sales_group TEXT NOT NULL,
          lead_type TEXT NOT NULL,
          leads INTEGER,
          revenue REAL,
          dpl REAL,
          source_record_ref TEXT NOT NULL,
          UNIQUE(snapshot_id, sales_group, lead_type)
        );
        CREATE INDEX IF NOT EXISTS idx_inside_sales_dpl_lookup
          ON inside_sales_dpl_fact(lead_type, observed_at, sales_group);

        CREATE TABLE IF NOT EXISTS collections_weekly_fact (
          fact_id INTEGER PRIMARY KEY,
          snapshot_id TEXT NOT NULL,
          source_id TEXT NOT NULL,
          observed_at TEXT NOT NULL,
          period_name TEXT NOT NULL,
          speaker TEXT NOT NULL,
          week_label TEXT NOT NULL,
          amount_into_collections REAL,
          weekly_collected REAL,
          source_record_ref TEXT NOT NULL,
          UNIQUE(snapshot_id, period_name, speaker, week_label)
        );
        CREATE INDEX IF NOT EXISTS idx_collections_weekly_lookup
          ON collections_weekly_fact(period_name, week_label, speaker);

        CREATE TABLE IF NOT EXISTS collections_pending_event_fact (
          fact_id INTEGER PRIMARY KEY,
          snapshot_id TEXT NOT NULL,
          source_id TEXT NOT NULL,
          observed_at TEXT NOT NULL,
          period_name TEXT NOT NULL,
          speaker TEXT NOT NULL,
          event_name TEXT NOT NULL,
          event_date TEXT,
          amount REAL,
          source_record_ref TEXT NOT NULL,
          UNIQUE(snapshot_id, period_name, speaker, event_name, event_date)
        );
        CREATE INDEX IF NOT EXISTS idx_collections_pending_lookup
          ON collections_pending_event_fact(period_name, event_date, speaker);

        DROP VIEW IF EXISTS latest_inside_sales_dpl;
        CREATE VIEW latest_inside_sales_dpl AS
        SELECT f.*
        FROM inside_sales_dpl_fact f
        JOIN (
          SELECT source_id, MAX(observed_at) AS observed_at
          FROM inside_sales_dpl_fact GROUP BY source_id
        ) latest
          ON latest.source_id=f.source_id AND latest.observed_at=f.observed_at;

        DROP VIEW IF EXISTS latest_collections_weekly;
        CREATE VIEW latest_collections_weekly AS
        SELECT f.*
        FROM collections_weekly_fact f
        JOIN (
          SELECT source_id, MAX(observed_at) AS observed_at
          FROM collections_weekly_fact GROUP BY source_id
        ) latest
          ON latest.source_id=f.source_id AND latest.observed_at=f.observed_at;
        """
    )


def upsert_registry(conn: sqlite3.Connection, spec: SourceSpec, checked_at: str) -> None:
    conn.execute(
        """
        INSERT INTO source_registry (
          source_id, source_name, source_type, source_family, trust_role,
          business_area, canonical_ref, local_path, enabled, access_status,
          owner_hint, notes, created_at, updated_at
        ) VALUES (?, ?, ?, ?, 'operational', ?, ?, ?, 1, 'ok', 'Harlow accountable; Sean execution',
                  'Hourly structured archive for dashboard and ad-hoc analytics.', ?, ?)
        ON CONFLICT(source_id) DO UPDATE SET
          source_name=excluded.source_name,
          source_type=excluded.source_type,
          source_family=excluded.source_family,
          business_area=excluded.business_area,
          canonical_ref=excluded.canonical_ref,
          local_path=excluded.local_path,
          enabled=1,
          access_status='ok',
          updated_at=excluded.updated_at
        """,
        (
            spec.source_id,
            spec.source_name,
            spec.source_type,
            spec.source_family,
            spec.business_area,
            spec.canonical_ref,
            str(spec.local_path),
            checked_at,
            checked_at,
        ),
    )


def get_or_create_snapshot(
    conn: sqlite3.Connection,
    spec: SourceSpec,
    content_hash: str,
    byte_size: int,
    checked_at: str,
    sheet_count: int | None,
    metadata: dict[str, Any],
) -> tuple[str, bool]:
    existing = conn.execute(
        "SELECT snapshot_id FROM source_snapshots WHERE source_id=? AND content_sha256=?",
        (spec.source_id, content_hash),
    ).fetchone()
    if existing:
        return str(existing[0]), False
    snapshot_id = f"{spec.source_id}:{content_hash[:24]}"
    conn.execute(
        """
        INSERT INTO source_snapshots (
          snapshot_id, source_id, snapshot_at, fetched_at, local_path,
          content_sha256, byte_size, workbook_sheet_count, raw_metadata_json
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            snapshot_id,
            spec.source_id,
            checked_at,
            checked_at,
            str(spec.local_path),
            content_hash,
            byte_size,
            sheet_count,
            json_text(metadata),
        ),
    )
    return snapshot_id, True


def iter_nonempty_rows(ws: Any) -> Iterable[tuple[int, list[Any]]]:
    for row_number, values in enumerate(ws.iter_rows(values_only=True), start=1):
        normalized = list(values)
        while normalized and normalized[-1] in (None, ""):
            normalized.pop()
        if any(value not in (None, "") for value in normalized):
            yield row_number, normalized


def xlsx_semantic_hash(path: Path) -> tuple[str, int, int]:
    """Hash data-only workbook contents, not volatile ZIP metadata.

    Google XLSX exports can have a new byte hash on every download even when
    every cell is unchanged. A byte hash would duplicate ~14k raw rows hourly.
    """
    digest = hashlib.sha256()
    workbook = load_workbook(path, read_only=True, data_only=True)
    total_rows = 0
    try:
        for ws in workbook.worksheets:
            digest.update(json_text({"sheet": ws.title}).encode("utf-8"))
            for row_number, values in iter_nonempty_rows(ws):
                total_rows += 1
                digest.update(json_text([row_number, values]).encode("utf-8"))
    finally:
        sheet_count = len(workbook.sheetnames)
        workbook.close()
    return digest.hexdigest(), sheet_count, total_rows


def archive_xlsx(
    conn: sqlite3.Connection,
    spec: SourceSpec,
    snapshot_id: str,
    checked_at: str,
    is_new: bool,
) -> tuple[int, int, str]:
    if not is_new:
        sheets = conn.execute(
            "SELECT COUNT(*) FROM workbook_sheets_raw WHERE snapshot_id=?",
            (snapshot_id,),
        ).fetchone()[0]
        rows = conn.execute(
            "SELECT COUNT(*) FROM source_rows_raw WHERE snapshot_id=? AND record_type='workbook_row'",
            (snapshot_id,),
        ).fetchone()[0]
        fingerprint = conn.execute(
            "SELECT COALESCE(MAX(schema_fingerprint),'') FROM workbook_sheets_raw WHERE snapshot_id=?",
            (snapshot_id,),
        ).fetchone()[0]
        return int(sheets), int(rows), str(fingerprint or "")

    workbook = load_workbook(spec.local_path, read_only=True, data_only=True)
    sheet_fingerprints: list[str] = []
    total_rows = 0
    try:
        for sheet_index, ws in enumerate(workbook.worksheets):
            rows = list(iter_nonempty_rows(ws))
            total_rows += len(rows)
            max_row = max((row_number for row_number, _ in rows), default=0)
            max_col = max((len(values) for _, values in rows), default=0)
            header = rows[0][1] if rows else []
            schema_fingerprint = sha256_text(json_text({"sheet": ws.title, "header": header, "max_col": max_col}))
            sheet_fingerprints.append(schema_fingerprint)
            conn.execute(
                """
                INSERT OR REPLACE INTO workbook_sheets_raw (
                  snapshot_id, workbook_name, sheet_name, sheet_index,
                  nonempty_rows, max_row, max_col, header_row_json,
                  raw_cells_json_path, schema_fingerprint, parser_status, notes
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, NULL, ?, 'raw_rows_loaded', ?)
                """,
                (
                    snapshot_id,
                    spec.source_name,
                    ws.title,
                    sheet_index,
                    len(rows),
                    max_row,
                    max_col,
                    json_text(header),
                    schema_fingerprint,
                    "Cells stored queryably in source_rows_raw as data_only JSON arrays.",
                ),
            )
            for row_number, values in rows:
                row_json = json_text(values)
                conn.execute(
                    """
                    INSERT OR IGNORE INTO source_rows_raw (
                      snapshot_id, source_id, record_type, sheet_name, row_number,
                      record_key, row_json, row_sha256, ingested_at
                    ) VALUES (?, ?, 'workbook_row', ?, ?, NULL, ?, ?, ?)
                    """,
                    (
                        snapshot_id,
                        spec.source_id,
                        ws.title,
                        row_number,
                        row_json,
                        sha256_text(row_json),
                        checked_at,
                    ),
                )
    finally:
        workbook.close()
    combined = sha256_text("|".join(sheet_fingerprints)) if sheet_fingerprints else ""
    return len(sheet_fingerprints), total_rows, combined


def archive_json_raw(
    conn: sqlite3.Connection,
    spec: SourceSpec,
    snapshot_id: str,
    data: dict[str, Any],
    checked_at: str,
    is_new: bool,
) -> int:
    if not is_new:
        return int(
            conn.execute(
                "SELECT COUNT(*) FROM source_rows_raw WHERE snapshot_id=? AND record_type LIKE 'json_%'",
                (snapshot_id,),
            ).fetchone()[0]
        )
    count = 0
    for key, value in data.items():
        records = value if isinstance(value, list) else [value]
        for index, record in enumerate(records):
            row_json = json_text(record)
            conn.execute(
                """
                INSERT OR IGNORE INTO source_rows_raw (
                  snapshot_id, source_id, record_type, sheet_name, row_number,
                  record_key, row_json, row_sha256, ingested_at
                ) VALUES (?, ?, ?, NULL, ?, ?, ?, ?, ?)
                """,
                (
                    snapshot_id,
                    spec.source_id,
                    f"json_{key}",
                    index + 1,
                    f"{key}:{index}",
                    row_json,
                    sha256_text(row_json),
                    checked_at,
                ),
            )
            count += 1
    return count


def normalize_inside_sales(
    conn: sqlite3.Connection,
    spec: SourceSpec,
    snapshot_id: str,
    data: dict[str, Any],
    checked_at: str,
) -> int:
    observed_at = str(data.get("lastUpdated") or checked_at)
    count = 0
    for group_index, group in enumerate(data.get("groups") or []):
        sales_group = str(group.get("group") or "Unknown")
        lead_rows = group.get("byLeadType") or []
        if not lead_rows:
            lead_rows = [{"leadType": "ALL", "leads": group.get("totalLeads"), "revenue": group.get("totalRevenue"), "dpl": group.get("ytdDpl")}]
        for lead_index, row in enumerate(lead_rows):
            lead_type = str(row.get("leadType") or "Unknown")
            conn.execute(
                """
                INSERT OR REPLACE INTO inside_sales_dpl_fact (
                  snapshot_id, source_id, observed_at, sales_group, lead_type,
                  leads, revenue, dpl, source_record_ref
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    snapshot_id,
                    spec.source_id,
                    observed_at,
                    sales_group,
                    lead_type,
                    row.get("leads"),
                    row.get("revenue"),
                    row.get("dpl"),
                    f"groups[{group_index}].byLeadType[{lead_index}]",
                ),
            )
            count += 1
    return count


def normalize_collections(
    conn: sqlite3.Connection,
    spec: SourceSpec,
    snapshot_id: str,
    data: dict[str, Any],
    checked_at: str,
) -> tuple[int, int]:
    weekly_count = 0
    pending_count = 0
    for period_name in ("past6Weeks", "past10Weeks", "ytd"):
        for speaker_index, speaker_row in enumerate(data.get(period_name) or []):
            speaker = str(speaker_row.get("speaker") or "Unknown")
            for week_index, week in enumerate(speaker_row.get("byWeek") or []):
                conn.execute(
                    """
                    INSERT OR REPLACE INTO collections_weekly_fact (
                      snapshot_id, source_id, observed_at, period_name, speaker,
                      week_label, amount_into_collections, weekly_collected,
                      source_record_ref
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        snapshot_id,
                        spec.source_id,
                        checked_at,
                        period_name,
                        speaker,
                        str(week.get("week") or "Unknown"),
                        week.get("amountIntoCollections"),
                        week.get("weeklyCollected"),
                        f"{period_name}[{speaker_index}].byWeek[{week_index}]",
                    ),
                )
                weekly_count += 1
    for period_name in ("past6WeeksPending", "past10WeeksPending", "ytdPending"):
        for speaker_index, speaker_row in enumerate(data.get(period_name) or []):
            speaker = str(speaker_row.get("speaker") or "Unknown")
            for event_index, event in enumerate(speaker_row.get("byEvent") or []):
                conn.execute(
                    """
                    INSERT OR REPLACE INTO collections_pending_event_fact (
                      snapshot_id, source_id, observed_at, period_name, speaker,
                      event_name, event_date, amount, source_record_ref
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        snapshot_id,
                        spec.source_id,
                        checked_at,
                        period_name,
                        speaker,
                        str(event.get("event") or "Unknown"),
                        event.get("date"),
                        event.get("amount"),
                        f"{period_name}[{speaker_index}].byEvent[{event_index}]",
                    ),
                )
                pending_count += 1
    return weekly_count, pending_count


def record_check(
    conn: sqlite3.Connection,
    spec: SourceSpec,
    checked_at: str,
    content_hash: str,
    byte_size: int,
    rows_seen: int,
    sheets_seen: int | None,
    schema_fingerprint: str,
    notes: str,
) -> None:
    conn.execute(
        """
        INSERT INTO source_check_runs (
          source_id, checked_at, status, transport_status, api_status,
          bytes_seen, rows_seen, sheets_seen, records_seen,
          latest_source_record_at, content_sha256, schema_fingerprint,
          reminder_needed, notes
        ) VALUES (?, ?, 'ok_fresh', 'local_export_ok', 'parsed', ?, ?, ?, ?, ?, ?, ?, 0, ?)
        """,
        (
            spec.source_id,
            checked_at,
            byte_size,
            rows_seen,
            sheets_seen,
            rows_seen,
            checked_at,
            content_hash,
            schema_fingerprint,
            notes,
        ),
    )


def archive_source(conn: sqlite3.Connection, spec: SourceSpec, checked_at: str) -> dict[str, Any]:
    if not spec.local_path.exists():
        raise FileNotFoundError(spec.local_path)
    raw = spec.local_path.read_bytes()
    file_hash = sha256_bytes(raw)
    upsert_registry(conn, spec, checked_at)

    if spec.kind == "xlsx":
        # Google exports rewrite ZIP metadata even when the workbook's cells do
        # not change. Use a semantic cell hash as the snapshot identity so the
        # warehouse grows only when queryable source truth changes.
        content_hash, sheet_count, semantic_rows = xlsx_semantic_hash(spec.local_path)
        snapshot_id, is_new = get_or_create_snapshot(
            conn,
            spec,
            content_hash,
            len(raw),
            checked_at,
            sheet_count,
            {
                "kind": "xlsx",
                "mode": "data_only",
                "loader": "archive_tlwb_structured_sources.py",
                "file_sha256": file_hash,
                "semantic_rows": semantic_rows,
            },
        )
        sheets_seen, rows_seen, schema_fingerprint = archive_xlsx(conn, spec, snapshot_id, checked_at, is_new)
        detail = f"{sheets_seen} sheets / {rows_seen} non-empty rows"
    else:
        content_hash = file_hash
        data = json.loads(raw)
        snapshot_id, is_new = get_or_create_snapshot(
            conn,
            spec,
            content_hash,
            len(raw),
            checked_at,
            None,
            {"kind": "json", "top_level_keys": list(data), "loader": "archive_tlwb_structured_sources.py"},
        )
        rows_seen = archive_json_raw(conn, spec, snapshot_id, data, checked_at, is_new)
        schema_fingerprint = sha256_text(json_text({key: type(value).__name__ for key, value in data.items()}))
        if spec.kind == "inside_sales_json":
            normalized = normalize_inside_sales(conn, spec, snapshot_id, data, checked_at)
            detail = f"{rows_seen} raw records / {normalized} normalized DPL facts"
        elif spec.kind == "collections_json":
            weekly, pending = normalize_collections(conn, spec, snapshot_id, data, checked_at)
            detail = f"{rows_seen} raw records / {weekly} weekly + {pending} pending-event facts"
        else:
            detail = f"{rows_seen} raw records"
        sheets_seen = None

    record_check(
        conn,
        spec,
        checked_at,
        content_hash,
        len(raw),
        rows_seen,
        sheets_seen,
        schema_fingerprint,
        f"Structured source archived; snapshot {'created' if is_new else 'reused'}; {detail}.",
    )
    return {
        "source_id": spec.source_id,
        "snapshot_id": snapshot_id,
        "new_snapshot": is_new,
        "detail": detail,
    }


def main() -> int:
    checked_at = now_iso()
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.execute("PRAGMA foreign_keys=ON")
    try:
        init_schema(conn)
        results = []
        with conn:
            for spec in SOURCES:
                results.append(archive_source(conn, spec, checked_at))
        quick = conn.execute("PRAGMA quick_check").fetchone()[0]
        if quick != "ok":
            raise RuntimeError(f"SQLite quick_check failed: {quick}")
        print(json.dumps({"checked_at": checked_at, "database": str(DB_PATH), "quick_check": quick, "sources": results}, indent=2))
        return 0
    finally:
        conn.close()


if __name__ == "__main__":
    raise SystemExit(main())
